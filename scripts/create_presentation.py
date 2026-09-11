from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "VendorLens_AI_20_Minute_Presentation.pptx"

NAVY = RGBColor(11, 31, 58)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(20, 184, 166)
VIOLET = RGBColor(124, 58, 237)
ORANGE = RGBColor(249, 115, 22)
INK = RGBColor(20, 33, 50)
MUTED = RGBColor(92, 108, 128)
PALE = RGBColor(241, 246, 252)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(22, 163, 74)
RED = RGBColor(220, 38, 38)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def shape(slide, x, y, w, h, fill, rounded=False, line_color=None):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line_color or fill
    if rounded:
        s.adjustments[0] = 0.12
    return s


def add_text(slide, value, x, y, w, h, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=15, color=INK):
    return add_text(slide, "\n".join("• " + item for item in items), x, y, w, h,
                    size=size, color=color)


def footer(slide, number, timing):
    shape(slide, 0, 7.1, 13.333, 0.4, NAVY)
    add_text(slide, "VendorLens AI", 0.65, 7.18, 2.2, 0.16, 8.5, RGBColor(195, 214, 234), True)
    add_text(slide, timing, 4.6, 7.18, 4.1, 0.16, 8.5, RGBColor(195, 214, 234), False, PP_ALIGN.CENTER)
    add_text(slide, f"{number:02d}", 12.1, 7.18, 0.55, 0.16, 9, RGBColor(136, 190, 255), True, PP_ALIGN.RIGHT)


def title(slide, kicker, heading, subtitle, number, timing):
    shape(slide, 0, 0, 13.333, 7.5, WHITE)
    shape(slide, 0, 0, 13.333, 0.14, BLUE)
    add_text(slide, kicker.upper(), 0.7, 0.45, 3.2, 0.22, 10.5, BLUE, True)
    add_text(slide, heading, 0.7, 0.78, 11.7, 0.58, 29, NAVY, True, font="Aptos Display")
    add_text(slide, subtitle, 0.72, 1.48, 11.6, 0.3, 13.5, MUTED)
    footer(slide, number, timing)


def clear_slide(slide):
    """Remove generated content when a slide needs a final client-facing variant."""
    for shape_obj in list(slide.shapes):
        shape_obj._element.getparent().remove(shape_obj._element)


# Slide 1: title
slide = prs.slides.add_slide(BLANK)
shape(slide, 0, 0, 13.333, 7.5, NAVY)
shape(slide, 0, 0, 13.333, 0.16, ORANGE)
shape(slide, 9.0, 0.16, 4.333, 7.34, RGBColor(15, 43, 76))
add_text(slide, "VendorLens AI", 0.82, 1.2, 7.2, 0.65, 40, WHITE, True, font="Aptos Display")
add_text(slide, "From supplier documents to an\nauditable onboarding decision", 0.86, 2.12, 7.6, 1.05, 27, RGBColor(215, 229, 244), True, font="Aptos Display")
add_text(slide, "AI-assisted supplier onboarding and compliance review", 0.88, 3.55, 7.3, 0.28, 16, ORANGE, True)
add_text(slide, "React • FastAPI • PostgreSQL • ChromaDB • Azure OpenAI", 0.88, 5.78, 7.6, 0.22, 11, RGBColor(174, 201, 228))
add_text(slide, "Human confirmation remains mandatory", 0.88, 6.22, 7.6, 0.22, 10.5, RGBColor(174, 201, 228))
for i, (label, color) in enumerate([("INTAKE", BLUE), ("RAG", TEAL), ("REVIEW", VIOLET), ("OBSERVE", ORANGE)]):
    y = 1.25 + i * 1.25
    shape(slide, 9.55, y, 2.85, 0.66, color, True)
    add_text(slide, label, 9.55, y + 0.18, 2.85, 0.22, 16, WHITE, True, PP_ALIGN.CENTER)
footer(slide, 1, "Opening | 0:00-0:45")


# Slide 2: problem and solution
slide = prs.slides.add_slide(BLANK)
title(slide, "Problem + solution", "Reduce review effort without giving AI the final vote",
      "The application automates repetitive understanding while preserving evidence, rules, and human accountability.", 2, "Problem & business case | 0:45-4:00")
shape(slide, 0.78, 2.15, 5.45, 4.4, NAVY, True)
add_text(slide, "THE PROBLEM", 1.1, 2.5, 2.0, 0.22, 11, ORANGE, True)
add_text(slide, "Supplier onboarding is a document bottleneck", 1.1, 2.88, 4.4, 0.62, 23, WHITE, True, font="Aptos Display")
bullets(slide, [
    "facts are spread across registration, tax, and insurance files",
    "follow-up questions require manual searching",
    "PII can leak into prompts or ad-hoc logs",
    "completeness and expiry checks are easy to miss",
], 1.1, 3.78, 4.45, 1.9, 15, RGBColor(220, 235, 249))
shape(slide, 6.55, 2.15, 5.95, 4.4, PALE, True, RGBColor(220, 230, 241))
add_text(slide, "THE SOLUTION", 6.9, 2.5, 2.0, 0.22, 11, BLUE, True)
add_text(slide, "One controlled workflow", 6.9, 2.88, 4.3, 0.42, 23, NAVY, True, font="Aptos Display")
flow = [
    ("1", "Intake", "Upload three required documents", BLUE),
    ("2", "AI", "Redact, extract, embed, retrieve", TEAL),
    ("3", "Review", "Citations + deterministic checks", VIOLET),
    ("4", "Decision", "Human approval → ERP handoff", ORANGE),
]
for i, (num, head, desc, accent) in enumerate(flow):
    y = 3.58 + i * 0.62
    shape(slide, 6.9, y, 0.38, 0.38, accent, True)
    add_text(slide, num, 6.9, y + 0.09, 0.38, 0.16, 11, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, head, 7.5, y + 0.01, 1.0, 0.2, 14, NAVY, True)
    add_text(slide, desc, 8.45, y + 0.01, 3.4, 0.2, 11.5, MUTED)
add_text(slide, "AI suggests → rules explain → reviewer confirms", 6.9, 6.05, 5.0, 0.24, 15, BLUE, True)


# Slide 3: architecture and tools
slide = prs.slides.add_slide(BLANK)
title(slide, "Architecture + tools", "A simple request path with an explicit control plane",
      "The application path is synchronous; evaluation, tracing, and monitoring observe the same running system.", 3, "Architecture & tools | 4:00-8:00")
# architecture flow
boxes = [
    (0.78, "React + MUI", "dashboard • intake • review", BLUE),
    (3.25, "FastAPI", "REST • validation • audit", TEAL),
    (5.75, "AI services", "redact • extract • retrieve", NAVY),
    (8.35, "Azure OpenAI", "gpt-4o-mini • embeddings", ORANGE),
    (10.85, "ChromaDB", "supplier-filtered vectors", VIOLET),
]
for i, (x, head, desc, accent) in enumerate(boxes):
    shape(slide, x, 2.05, 2.0, 0.95, accent, True)
    add_text(slide, head, x, 2.27, 2.0, 0.23, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, desc, x + 0.08, 2.62, 1.84, 0.16, 8.8, RGBColor(226, 239, 252), False, PP_ALIGN.CENTER)
    if i < len(boxes) - 1:
        add_text(slide, "→", x + 2.06, 2.32, 0.35, 0.25, 20, accent, True, PP_ALIGN.CENTER)
shape(slide, 2.2, 3.55, 3.95, 1.05, PALE, True, RGBColor(220, 230, 241))
add_text(slide, "PostgreSQL", 2.48, 3.8, 1.5, 0.22, 17, NAVY, True)
add_text(slide, "suppliers • documents • AI runs • compliance • audit", 2.48, 4.16, 3.3, 0.18, 10.5, MUTED)
shape(slide, 7.0, 3.55, 4.15, 1.05, PALE, True, RGBColor(220, 230, 241))
add_text(slide, "Quality + operations", 7.28, 3.8, 2.2, 0.22, 17, NAVY, True)
add_text(slide, "Langfuse • Promptfoo • Prometheus", 7.28, 4.16, 3.0, 0.18, 10.5, MUTED)
add_text(slide, "Tool choices", 0.82, 5.15, 1.45, 0.23, 16, NAVY, True)
add_text(slide, "ChromaDB = low-permission deployment  |  gpt-4o-mini = cost/latency  |  rules = explainable decisions", 2.45, 5.15, 9.7, 0.24, 13, MUTED)
add_text(slide, "Privacy: PII is redacted before applicable LLM calls; telemetry labels never contain supplier data.", 0.82, 5.92, 11.6, 0.25, 14, ORANGE, True)

# Replace the draft architecture with a single readable system flow.
clear_slide(slide)
title(slide, "Architecture + tools", "From intake to an auditable decision",
      "The same controlled path supports document processing, supplier Q&A, compliance, and operations.", 3, "Architecture & tools | 4:00-8:00")

flow_boxes = [
    (0.75, 2.15, 2.2, "React + MUI", "Intake • review • Q&A", BLUE),
    (3.45, 2.15, 2.2, "FastAPI", "API • validation • audit", TEAL),
    (6.15, 2.15, 2.45, "Workflow services", "redact • extract • retrieve", NAVY),
    (9.2, 2.15, 2.45, "Azure OpenAI", "GPT-4o-mini • embeddings", ORANGE),
]
for index, (x, y, width, head, desc, accent) in enumerate(flow_boxes):
    shape(slide, x, y, width, 0.86, accent, True)
    add_text(slide, head, x, y + 0.18, width, 0.22, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, desc, x + 0.08, y + 0.5, width - 0.16, 0.15, 9, RGBColor(226, 239, 252), False, PP_ALIGN.CENTER)
    if index < len(flow_boxes) - 1:
        add_text(slide, "→", x + width + 0.2, y + 0.27, 0.35, 0.25, 20, accent, True, PP_ALIGN.CENTER)

add_text(slide, "↓", 4.35, 3.12, 0.3, 0.25, 20, TEAL, True, PP_ALIGN.CENTER)
add_text(slide, "↓", 6.95, 3.12, 0.3, 0.25, 20, NAVY, True, PP_ALIGN.CENTER)
add_text(slide, "↓", 10.2, 3.12, 0.3, 0.25, 20, ORANGE, True, PP_ALIGN.CENTER)

branches = [
    (1.1, 3.65, 3.0, "PostgreSQL", "supplier records • AI runs • compliance • audit", BLUE),
    (4.85, 3.65, 3.0, "ChromaDB", "supplier-scoped chunks and embeddings", VIOLET),
    (8.6, 3.65, 3.0, "Observability", "Langfuse • Promptfoo • Prometheus", ORANGE),
]
for x, y, width, head, desc, accent in branches:
    shape(slide, x, y, width, 0.92, PALE, True, accent)
    add_text(slide, head, x + 0.12, y + 0.18, width - 0.24, 0.22, 15, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, desc, x + 0.12, y + 0.52, width - 0.24, 0.16, 9.5, MUTED, False, PP_ALIGN.CENTER)

add_text(slide, "↓", 2.55, 4.82, 0.3, 0.25, 20, BLUE, True, PP_ALIGN.CENTER)
add_text(slide, "↓", 6.3, 4.82, 0.3, 0.25, 20, VIOLET, True, PP_ALIGN.CENTER)
shape(slide, 2.35, 5.25, 3.6, 0.82, NAVY, True)
add_text(slide, "Compliance rules + audit trail", 2.35, 5.5, 3.6, 0.22, 14, WHITE, True, PP_ALIGN.CENTER)
shape(slide, 7.2, 5.25, 3.8, 0.82, GREEN, True)
add_text(slide, "Human decision → ERP handoff", 7.2, 5.5, 3.8, 0.22, 14, WHITE, True, PP_ALIGN.CENTER)
add_text(slide, "→", 6.25, 5.49, 0.45, 0.25, 20, ORANGE, True, PP_ALIGN.CENTER)
add_text(slide, "AI suggests • rules explain • reviewer confirms", 3.25, 6.35, 6.9, 0.22, 14, BLUE, True, PP_ALIGN.CENTER)


# Slide 4: demo
slide = prs.slides.add_slide(BLANK)
title(slide, "Live demo", "One supplier, one complete decision path",
      "Use the synthetic Kaveri Flow Controls pack. Keep the browser, backend terminal, and Prometheus UI ready.", 4, "Live demo | 8:00-15:00")
demo = [
    ("01", "Create", "Supplier name, country, optional email", BLUE),
    ("02", "Upload", "Registration + tax + insurance; show replace/delete", TEAL),
    ("03", "Process", "PII redaction → fields → chunks", VIOLET),
    ("04", "Ask", "Cited answer + absent bank balance → not found", ORANGE),
    ("05", "Review", "Correct one field; rerun compliance", NAVY),
    ("06", "Decide", "Approve → mock ERP ID + audit event", GREEN),
]
for i, (num, head, desc, accent) in enumerate(demo):
    y = 2.05 + i * 0.67
    shape(slide, 0.95, y, 0.52, 0.46, accent, True)
    add_text(slide, num, 0.95, y + 0.13, 0.52, 0.16, 11, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, head, 1.85, y + 0.1, 1.25, 0.22, 16, NAVY, True)
    add_text(slide, desc, 3.35, y + 0.1, 7.55, 0.22, 13, MUTED)
    if i < 5:
        add_text(slide, "↓", 1.11, y + 0.47, 0.2, 0.18, 14, RGBColor(170, 187, 207), True, PP_ALIGN.CENTER)
shape(slide, 0.95, 6.25, 11.15, 0.45, RGBColor(239, 246, 255), True, RGBColor(191, 219, 254))
add_text(slide, "Fallback: use QUALITY_EVALUATION.md and latest results if an external AI call is slow.", 1.2, 6.38, 10.65, 0.18, 11.5, BLUE, True, PP_ALIGN.CENTER)

# Keep the hand-off slide intentionally minimal; the application is shown live.
clear_slide(slide)
shape(slide, 0, 0, 13.333, 7.5, PALE)
shape(slide, 0, 0, 13.333, 0.16, VIOLET)
add_text(slide, "Live demo", 0.82, 2.45, 11.7, 0.9, 54, NAVY, True, PP_ALIGN.CENTER, font="Aptos Display", valign=MSO_ANCHOR.MIDDLE)
footer(slide, 4, "Live walkthrough | 8:00-15:00")


# Slide 5: evaluation and monitoring
slide = prs.slides.add_slide(BLANK)
title(slide, "Evidence + observability", "The system is measurable, not just impressive",
      "Quality tells us whether answers are good; observability tells us what happened and whether the app is healthy.", 5, "Evidence & operations | 15:00-17:30")
cards = [
    ("100%", "field checks", "4 suppliers • 12 PDFs", BLUE),
    ("100%", "Q&A + citations", "controlled corpus", TEAL),
    ("100%", "not-found + isolation", "no cross-supplier answers", VIOLET),
    ("2.0s", "average Q&A latency", "2.4s maximum in report", ORANGE),
]
for i, (value, label, note, accent) in enumerate(cards):
    x = 0.82 + i * 3.05
    shape(slide, x, 2.1, 2.72, 1.23, PALE, True, RGBColor(220, 230, 241))
    shape(slide, x, 2.1, 0.08, 1.23, accent, True)
    add_text(slide, value, x + 0.23, 2.3, 2.2, 0.35, 25, NAVY, True, font="Aptos Display")
    add_text(slide, label, x + 0.23, 2.75, 2.25, 0.2, 12, NAVY, True)
    add_text(slide, note, x + 0.23, 3.03, 2.25, 0.16, 9.5, MUTED)
add_text(slide, "Three feedback loops", 0.84, 3.86, 2.4, 0.25, 17, NAVY, True)
loops = [
    ("Langfuse", "trace model, tokens, latency", BLUE),
    ("Promptfoo", "assert citations, boundaries, not-found", TEAL),
    ("Prometheus", "scrape traffic, errors, p95, dependencies", ORANGE),
]
for i, (name, desc, accent) in enumerate(loops):
    x = 0.9 + i * 4.1
    shape(slide, x, 4.35, 3.55, 1.15, WHITE, True, accent)
    add_text(slide, name, x + 0.25, 4.62, 1.3, 0.22, 16, accent, True)
    add_text(slide, desc, x + 1.5, 4.59, 1.72, 0.34, 11, MUTED)
add_text(slide, "Demo caveat: these are controlled synthetic evaluations, not production guarantees.", 0.9, 6.1, 11.2, 0.25, 13.5, ORANGE, True, PP_ALIGN.CENTER)

# Results will be appended after the validation run; show only the measurement approach now.
clear_slide(slide)
title(slide, "Quality + observability", "Built-in measurement for quality and operations",
      "The application captures the evidence needed to validate answer quality and service health.", 5, "Quality & operations | 15:00-17:30")
cards = [
    ("Langfuse", "Trace model calls, tokens, latency, and failures", BLUE),
    ("Promptfoo", "Evaluate grounding, citations, isolation, and refusals", TEAL),
    ("Prometheus", "Monitor traffic, errors, latency, and dependencies", ORANGE),
]
for i, (name, desc, accent) in enumerate(cards):
    x = 0.9 + i * 4.1
    shape(slide, x, 2.2, 3.55, 1.45, WHITE, True, accent)
    add_text(slide, name, x + 0.25, 2.52, 3.0, 0.25, 18, accent, True)
    add_text(slide, desc, x + 0.25, 2.95, 3.0, 0.42, 12, MUTED)
add_text(slide, "Validation results will be added after the monitoring and quality test run.", 0.95, 5.1, 11.4, 0.35, 18, NAVY, True, PP_ALIGN.CENTER)


# Slide 6: tradeoffs, limitations, future
slide = prs.slides.add_slide(BLANK)
title(slide, "Trade-offs + future scope", "Pragmatic choices today; production hardening next",
      "The important design question is not whether AI can answer—it is where AI must stop.", 6, "Trade-offs & future | 17:30-19:15")
shape(slide, 0.82, 2.1, 5.65, 4.3, RGBColor(255, 247, 237), True, RGBColor(254, 215, 170))
add_text(slide, "Trade-offs accepted", 1.15, 2.47, 2.5, 0.3, 20, ORANGE, True, font="Aptos Display")
bullets(slide, [
    "ChromaDB over pgvector: fewer permissions and faster setup",
    "gpt-4o-mini default: lower cost and latency",
    "rules + human review: explainability over autonomy",
    "Prometheus metrics: operational visibility with a lightweight deployment",
], 1.15, 3.12, 4.7, 2.25, 13.5, RGBColor(112, 61, 12))
shape(slide, 6.78, 2.1, 5.72, 4.3, RGBColor(239, 246, 255), True, RGBColor(191, 219, 254))
add_text(slide, "Limitations → future", 7.1, 2.47, 2.8, 0.3, 20, BLUE, True, font="Aptos Display")
bullets(slide, [
    "scanned PDFs → OCR and multi-page parsing",
    "no auth/RBAC → SSO, retention, encryption, audit controls",
    "small corpus → adversarial held-out evaluation",
    "ERP handoff → idempotent production integration",
    "metrics endpoint → hosted dashboards and alerting",
], 7.1, 3.12, 4.75, 2.25, 13.5, NAVY)
add_text(slide, "North star: reduce reviewer effort while keeping evidence and accountability visible.", 1.0, 6.62, 11.3, 0.23, 14, NAVY, True, PP_ALIGN.CENTER)


# Slide 7: Q&A
slide = prs.slides.add_slide(BLANK)
shape(slide, 0, 0, 13.333, 7.5, PALE)
shape(slide, 0, 0, 13.333, 0.16, VIOLET)
add_text(slide, "Q&A", 0.82, 1.15, 4.3, 0.75, 58, NAVY, True, font="Aptos Display")
add_text(slide, "Questions, challenges, and trade-offs", 0.88, 2.18, 7.8, 0.38, 24, MUTED, False, font="Aptos Display")
shape(slide, 0.9, 3.35, 11.4, 1.52, WHITE, True, RGBColor(220, 230, 241))
add_text(slide, "Discussion topics", 1.25, 3.7, 2.4, 0.23, 15, VIOLET, True)
add_text(slide, "Why ChromaDB?  •  How do you prove grounding?  •  What happens when extraction is wrong?  •  How would this reach production?", 3.7, 3.63, 7.9, 0.52, 14, NAVY, True)
add_text(slide, "Thank you", 0.92, 5.7, 2.4, 0.34, 22, NAVY, True, font="Aptos Display")
footer(slide, 7, "Questions & discussion")

# Keep the closing page focused only on Q&A.
clear_slide(slide)
shape(slide, 0, 0, 13.333, 7.5, PALE)
shape(slide, 0, 0, 13.333, 0.16, VIOLET)
add_text(slide, "Q&A", 0.82, 2.35, 11.7, 1.0, 62, NAVY, True, PP_ALIGN.CENTER, font="Aptos Display", valign=MSO_ANCHOR.MIDDLE)
footer(slide, 7, "Questions & discussion")

prs.core_properties.title = "VendorLens AI - 20 Minute Presentation"
prs.core_properties.subject = "Supplier onboarding, RAG, compliance review, and observability"
prs.core_properties.author = "VendorLens AI"
prs.save(OUT)
print(OUT)
